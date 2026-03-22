from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pandas as pd
from datetime import datetime

# 读取数据
sales_df = pd.read_excel('data/2023年8月-9月销售记录.xlsx')
sales_df['销售额'] = sales_df['单价(元)'] * sales_df['销售量']
sales_df['月份'] = pd.to_datetime(sales_df['销售日期']).dt.month

# 计算数据
august_sales = sales_df[sales_df['月份'] == 8]['销售额'].sum()
september_sales = sales_df[sales_df['月份'] == 9]['销售额'].sum()
product_sales = sales_df.groupby('产品名')['销售额'].sum().sort_values(ascending=False)
supplier_sales = sales_df.groupby('供应商')['销售额'].sum().sort_values(ascending=False)

# 创建Word文档
doc = Document()

# 标题
title = doc.add_heading('2023年8月-9月销售数据分析报告', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 报告日期
date_para = doc.add_paragraph()
date_str = datetime.now().strftime("%Y年%m月%d日")
date_para.add_run(f'报告生成日期: {date_str}').italic = True
date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()

# 一、概述
doc.add_heading('一、报告概述', 1)
doc.add_paragraph('本报告基于2023年8月至9月的销售记录数据，对销售额、产品销售情况及供应商表现进行分析，为业务决策提供数据支持。')

# 二、月度销售额对比
doc.add_heading('二、月度销售额对比', 1)

table1 = doc.add_table(rows=3, cols=3)
table1.style = 'Light Grid Accent 1'
hdr_cells = table1.rows[0].cells
hdr_cells[0].text = '月份'
hdr_cells[1].text = '销售额(元)'
hdr_cells[2].text = '占比'

total_sales = august_sales + september_sales
row1 = table1.rows[1].cells
row1[0].text = '8月'
row1[1].text = f'{august_sales:,.2f}'
row1[2].text = f'{august_sales/total_sales*100:.2f}%'

row2 = table1.rows[2].cells
row2[0].text = '9月'
row2[1].text = f'{september_sales:,.2f}'
row2[2].text = f'{september_sales/total_sales*100:.2f}%'

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run('分析结论: ').bold = True
growth_rate = (september_sales - august_sales) / august_sales * 100
p.add_run(f'9月销售额({september_sales:,.2f}元)较8月销售额({august_sales:,.2f}元)增长{growth_rate:.2f}%，整体呈上升趋势。')

# 三、产品销售分析
doc.add_heading('三、产品销售分析', 1)

p = doc.add_paragraph()
p.add_run('销售总额最大产品: ').bold = True
p.add_run(product_sales.index[0]).bold = True

doc.add_paragraph(f'销售额: {product_sales.iloc[0]:,.2f} 元')

doc.add_paragraph('产品销售TOP5:')
for i, (product, sales) in enumerate(product_sales.head(5).items(), 1):
    doc.add_paragraph(f'{i}. {product}: {sales:,.2f} 元', style='List Number')

# 四、供应商分析
doc.add_heading('四、供应商销售分析', 1)

p = doc.add_paragraph()
p.add_run('销售额最好供应商: ').bold = True
p.add_run(supplier_sales.index[0]).bold = True

doc.add_paragraph(f'销售额: {supplier_sales.iloc[0]:,.2f} 元')

doc.add_paragraph('供应商销售TOP5:')
for i, (supplier, sales) in enumerate(supplier_sales.head(5).items(), 1):
    doc.add_paragraph(f'{i}. {supplier}: {sales:,.2f} 元', style='List Number')

# 五、总结
doc.add_heading('五、总结与建议', 1)
doc.add_paragraph('1. 9月销售业绩优于8月，建议继续保持当前销售策略。')
doc.add_paragraph('2. PlayStation 5为销售冠军产品，可考虑增加库存和推广力度。')
doc.add_paragraph('3. 沈阳娱乐有限公司表现突出，可加强合作关系。')

# 保存文档
doc.save('销售数据分析报告.docx')
print('报告已生成: 销售数据分析报告.docx')
